import io
import re
from typing import Dict, Any, List, Optional
import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.services.ingestion.normalizer import (
    calculate_file_hash,
    normalize_section_title,
    extract_in_text_citations,
    parse_reference_entries,
)
from app.services.ingestion.ocr_detector import OCRDetector
from app.services.ingestion.document_validator import (
    DocumentPlausibilityValidator,
    DocumentValidationResult,
)


class ExtractedDocument:
    def __init__(
        self,
        title: str,
        file_hash: str,
        page_count: int,
        word_count: int,
        is_scanned: bool,
        ocr_required: bool,
        sections: List[Dict[str, Any]],
        references: List[Dict[str, Any]],
        citations: List[Dict[str, Any]],
        figures: List[Dict[str, Any]],
        tables: List[Dict[str, Any]],
        layout_metrics: Dict[str, Any],
        raw_text: str,
        validation: Optional[DocumentValidationResult] = None,
    ):
        self.title = title
        self.file_hash = file_hash
        self.page_count = page_count
        self.word_count = word_count
        self.is_scanned = is_scanned
        self.ocr_required = ocr_required
        self.sections = sections
        self.references = references
        self.citations = citations
        self.figures = figures
        self.tables = tables
        self.layout_metrics = layout_metrics
        self.raw_text = raw_text
        self.validation = validation or DocumentValidationResult()


class DocumentExtractor:
    @classmethod
    def extract_document(cls, content: bytes, filename: str) -> ExtractedDocument:
        suffix = filename.lower().split(".")[-1]
        if suffix == "pdf":
            return cls._extract_pdf(content, filename)
        elif suffix in ("docx", "doc"):
            return cls._extract_docx(content, filename)
        else:
            raise ValueError(f"Unsupported document format: .{suffix}. Supported formats are PDF and DOCX.")

    @classmethod
    def _extract_pdf(cls, content: bytes, filename: str) -> ExtractedDocument:
        file_hash = calculate_file_hash(content)
        is_scanned, ocr_required, ocr_reason = OCRDetector.inspect_pdf_scanned_state(content)

        doc = fitz.open(stream=content, filetype="pdf")
        page_count = len(doc)
        total_word_count = 0

        pages_data = []
        raw_text_full = []
        font_sizes = []
        margins = []

        figure_captions = []
        table_captions = []

        fig_pattern = re.compile(r"^(fig(?:ure|\.)?\s*(\d+|[A-Z]))[:\.\-\s]+(.*)", re.IGNORECASE)
        tbl_pattern = re.compile(r"^(table\s*(\d+|[IVXLCDM]+))[:\.\-\s]+(.*)", re.IGNORECASE)

        for page_idx in range(page_count):
            page = doc[page_idx]
            page_num = page_idx + 1
            rect = page.rect
            margins.append({"page": page_num, "width": rect.width, "height": rect.height})

            # Blocks extraction: (x0, y0, x1, y1, text, block_no, block_type)
            blocks = page.get_text("blocks")
            page_text = ""

            for b in blocks:
                # b[4] is text, b[6] is block_type (0 for text, 1 for image)
                if b[6] == 0:
                    block_text = b[4].strip()
                    if not block_text:
                        continue
                    page_text += block_text + "\n\n"

                    # Check for figure or table caption
                    lines = block_text.split("\n")
                    for line in lines:
                        line_str = line.strip()
                        fig_m = fig_pattern.match(line_str)
                        if fig_m:
                            figure_captions.append({
                                "figure_num": fig_m.group(2),
                                "caption": line_str,
                                "page": page_num,
                            })
                        tbl_m = tbl_pattern.match(line_str)
                        if tbl_m:
                            table_captions.append({
                                "table_num": tbl_m.group(2),
                                "caption": line_str,
                                "page": page_num,
                            })

            # Also check text spans for font sizes
            text_dict = page.get_text("dict")
            for block in text_dict.get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        if "size" in span and span.get("text", "").strip():
                            font_sizes.append(span["size"])

            pages_data.append({"page_num": page_num, "text": page_text})
            raw_text_full.append(page_text)
            total_word_count += len(page_text.split())

        doc.close()
        full_text = "\n\n".join(raw_text_full)

        # Detect sections across pages
        sections = cls._segment_sections_from_pages(pages_data)

        # Detect title (heuristically: first heading or first prominent line)
        title = filename
        if sections and sections[0]["normalized_title"] != "references":
            first_lines = [l.strip() for l in sections[0]["content"].split("\n") if len(l.strip()) > 5]
            if first_lines:
                title = first_lines[0][:200]

        # Extract in-text citations from body sections (excluding the references section itself)
        all_citations = []
        for s in sections:
            if s["normalized_title"] != "references":
                s_cits = extract_in_text_citations(s["content"], page_number=s.get("start_page", 1))
                all_citations.extend(s_cits)

        # Extract references
        references_text = ""
        for s in sections:
            if s["normalized_title"] == "references":
                references_text = s["content"]
                break

        references = parse_reference_entries(references_text)

        avg_font_size = sum(font_sizes) / len(font_sizes) if font_sizes else 11.0
        layout_metrics = {
            "avg_font_size_pt": round(avg_font_size, 2),
            "page_dimensions": margins[:3],
            "total_pages": page_count,
            "ocr_diagnostic": ocr_reason,
        }

        doc_validation = DocumentPlausibilityValidator.validate(
            raw_text=full_text,
            page_count=page_count,
            word_count=total_word_count,
            is_scanned=is_scanned,
            sections=sections,
        )

        return ExtractedDocument(
            title=title,
            file_hash=file_hash,
            page_count=page_count,
            word_count=total_word_count,
            is_scanned=is_scanned,
            ocr_required=ocr_required,
            sections=sections,
            references=references,
            citations=all_citations,
            figures=figure_captions,
            tables=table_captions,
            layout_metrics=layout_metrics,
            raw_text=full_text,
            validation=doc_validation,
        )

    @classmethod
    def _extract_docx(cls, content: bytes, filename: str) -> ExtractedDocument:
        file_hash = calculate_file_hash(content)
        doc = DocxDocument(io.BytesIO(content))

        sections_list = []
        current_section = {
            "section_order": 1,
            "title": "Header & Introduction",
            "normalized_title": "introduction",
            "content": "",
            "start_page": 1,
            "end_page": 1,
        }

        total_words = 0
        raw_paragraphs = []
        figures = []
        tables = []

        fig_pattern = re.compile(r"^(fig(?:ure|\.)?\s*(\d+|[A-Z]))[:\.\-\s]+(.*)", re.IGNORECASE)
        tbl_pattern = re.compile(r"^(table\s*(\d+|[IVXLCDM]+))[:\.\-\s]+(.*)", re.IGNORECASE)

        doc_title = None

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            raw_paragraphs.append(text)
            total_words += len(text.split())

            # Detect title from style
            if not doc_title and p.style and p.style.name.lower() in ("title", "heading 1"):
                doc_title = text[:200]

            # Detect headings
            is_heading = False
            if p.style and "heading" in p.style.name.lower():
                is_heading = True
            elif len(text) < 100 and (text.isupper() or re.match(r"^(\d+\.?|[IVXLCDM]+\.?)\s+[A-Z]", text)):
                is_heading = True

            fig_m = fig_pattern.match(text)
            if fig_m:
                figures.append({"figure_num": fig_m.group(2), "caption": text, "page": 1})

            tbl_m = tbl_pattern.match(text)
            if tbl_m:
                tables.append({"table_num": tbl_m.group(2), "caption": text, "page": 1})

            if is_heading:
                if current_section["content"].strip():
                    current_section["word_count"] = len(current_section["content"].split())
                    sections_list.append(current_section)
                current_section = {
                    "section_order": len(sections_list) + 1,
                    "title": text,
                    "normalized_title": normalize_section_title(text),
                    "content": "",
                    "start_page": 1,
                    "end_page": 1,
                }
            else:
                current_section["content"] += text + "\n\n"

        if current_section["content"].strip():
            current_section["word_count"] = len(current_section["content"].split())
            sections_list.append(current_section)

        # Process tables
        for idx, tbl in enumerate(doc.tables, 1):
            row_texts = []
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                row_texts.append(" | ".join(cells))
            tables.append({
                "table_num": str(idx),
                "caption": f"Embedded DOCX Table {idx}",
                "rows_count": len(tbl.rows),
                "page": 1,
            })

        full_text = "\n\n".join(raw_paragraphs)
        if not doc_title:
            doc_title = filename

        # Citations (excluding references section)
        citations = []
        for s in sections_list:
            if s["normalized_title"] != "references":
                citations.extend(extract_in_text_citations(s["content"], page_number=s.get("start_page", 1)))

        # References
        references_text = ""
        for s in sections_list:
            if s["normalized_title"] == "references":
                references_text = s["content"]
                break
        references = parse_reference_entries(references_text)

        estimated_pages = max(1, round(total_words / 500))

        layout_metrics = {
            "avg_font_size_pt": 12.0,
            "estimated_pages": estimated_pages,
            "format": "docx",
        }

        doc_validation = DocumentPlausibilityValidator.validate(
            raw_text=full_text,
            page_count=estimated_pages,
            word_count=total_words,
            is_scanned=False,
            sections=sections_list,
        )

        return ExtractedDocument(
            title=doc_title,
            file_hash=file_hash,
            page_count=estimated_pages,
            word_count=total_words,
            is_scanned=False,
            ocr_required=False,
            sections=sections_list,
            references=references,
            citations=citations,
            figures=figures,
            tables=tables,
            layout_metrics=layout_metrics,
            raw_text=full_text,
            validation=doc_validation,
        )

    @staticmethod
    def _segment_sections_from_pages(pages_data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        sections = []
        heading_regex = re.compile(
            r"^(?:(?:\d+|[IVXLCDM]+)[\.\:\-\s]+)?\s*(abstract|introduction|related\s+work|background|methodology|methods|proposed\s+method|experiments|experimental\s+results|results|discussion|conclusion|conclusions|references|bibliography)\b",
            re.IGNORECASE,
        )

        current_section = {
            "section_order": 1,
            "title": "Document Header & Preamble",
            "normalized_title": "preamble",
            "content": "",
            "start_page": 1,
            "end_page": 1,
        }

        for page in pages_data:
            p_num = page["page_num"]
            lines = page["text"].split("\n")

            for line in lines:
                l_str = line.strip()
                if not l_str:
                    continue

                if len(l_str) < 80 and heading_regex.match(l_str):
                    if current_section["content"].strip():
                        current_section["word_count"] = len(current_section["content"].split())
                        current_section["end_page"] = p_num
                        sections.append(current_section)

                    current_section = {
                        "section_order": len(sections) + 1,
                        "title": l_str,
                        "normalized_title": normalize_section_title(l_str),
                        "content": "",
                        "start_page": p_num,
                        "end_page": p_num,
                    }
                else:
                    current_section["content"] += l_str + "\n"

        if current_section["content"].strip():
            current_section["word_count"] = len(current_section["content"].split())
            current_section["end_page"] = pages_data[-1]["page_num"] if pages_data else 1
            sections.append(current_section)

        return sections
