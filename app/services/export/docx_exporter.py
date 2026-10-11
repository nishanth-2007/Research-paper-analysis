import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.schemas.report import CompleteAnalysisReport


class DOCXReportExporter:
    @classmethod
    def generate_docx_bytes(cls, report: CompleteAnalysisReport) -> bytes:
        doc = Document()

        # Title
        title_p = doc.add_paragraph()
        run = title_p.add_run("Journal Readiness & Verification Report")
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(26, 54, 93)

        # Manuscript Details
        doc.add_paragraph(f"Manuscript Title: {report.manuscript_title}")
        if report.target_journal_name:
            doc.add_paragraph(f"Target Journal: {report.target_journal_name} ({report.guideline_version_tag or 'Latest'})")
        doc.add_paragraph(f"Report Generated: {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')}")

        # Overall Score
        score_p = doc.add_paragraph()
        score_run = score_p.add_run(f"Readiness Score: {report.readiness_score.overall_score:.1f} / 100 ({report.readiness_score.readiness_tier})")
        score_run.bold = True
        score_run.font.size = Pt(14)

        # Executive Summary
        h2 = doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(report.executive_summary.synthesis)

        # Category Breakdown Table
        doc.add_heading("Category Breakdown", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Category"
        hdr_cells[1].text = "Weight"
        hdr_cells[2].text = "Score"
        hdr_cells[3].text = "Findings"

        for cat_id, cat_data in report.readiness_score.category_scores.items():
            row_cells = table.add_row().cells
            row_cells[0].text = cat_data.category_label
            row_cells[1].text = f"{cat_data.weight * 100:.0f}%"
            row_cells[2].text = f"{cat_data.score:.1f} / 100"
            row_cells[3].text = f"{cat_data.failed_count} fails, {cat_data.warning_count} warnings"

        # Prioritized Action Plan
        if report.consolidated_issues:
            doc.add_heading("Prioritized Action Plan", level=2)
            blockers = [c for c in report.consolidated_issues if c.tier.value == "submission_blocker"]
            improvements = [c for c in report.consolidated_issues if c.tier.value == "improvement_recommended"]

            if blockers:
                doc.add_heading("Submission Blockers", level=3)
                for b in blockers:
                    p = doc.add_paragraph(style="List Bullet")
                    r_title = p.add_run(f"[{b.severity.upper()}] {b.title}: ")
                    r_title.bold = True
                    p.add_run(f"{b.why_it_matters} Action: {b.recommended_action}")

            if improvements:
                doc.add_heading("Recommended Improvements", level=3)
                for imp in improvements:
                    p = doc.add_paragraph(style="List Bullet")
                    r_title = p.add_run(f"[{imp.severity.upper()}] {imp.title}: ")
                    r_title.bold = True
                    p.add_run(f"{imp.why_it_matters} Action: {imp.recommended_action}")

        elif report.actionable_recommendations:
            doc.add_heading("Priority Action Items", level=2)
            for rec in report.actionable_recommendations:
                p = doc.add_paragraph(style="List Bullet")
                r_title = p.add_run(f"[{rec.severity.upper()}] {rec.title}: ")
                r_title.bold = True
                p.add_run(rec.action_item)

        # Disclaimer
        doc.add_heading("Disclaimer", level=2)
        disc_p = doc.add_paragraph()
        disc_run = disc_p.add_run(report.readiness_score.disclaimer)
        disc_run.italic = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
