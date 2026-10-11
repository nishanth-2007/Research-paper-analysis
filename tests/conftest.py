import io
import os
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
import fitz
from docx import Document as DocxDoc

from app.database import Base, get_db
from app.models.user import User
from app.models.journal import Journal, GuidelineVersion, JournalRule, GuidelineChunk
from app.api.deps import DEFAULT_DEMO_USER_ID, DEFAULT_DEMO_EMAIL
from app.services.engines.rag.knowledge_base import DEFAULT_JOURNAL_SEEDS, prepare_seed_embeddings
from app.services.ingestion.extractor import ExtractedDocument
from main import app

TEST_DB_URL = "sqlite:///./test_platform.db"
test_engine = create_engine(TEST_DB_URL, connect_args={"check_same_thread": False})
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(scope="session", autouse=True)
def setup_test_db():
    Base.metadata.create_all(bind=test_engine)
    db = TestingSessionLocal()
    # Ensure demo user exists
    user = db.query(User).filter(User.id == DEFAULT_DEMO_USER_ID).first()
    if not user:
        user = User(
            id=DEFAULT_DEMO_USER_ID,
            email=DEFAULT_DEMO_EMAIL,
            hashed_password="hashed_demo_pw",
            full_name="Dr. Jane Researcher",
            is_active=True,
        )
        db.add(user)

    # Seed journals
    seeds = prepare_seed_embeddings(DEFAULT_JOURNAL_SEEDS)
    for s in seeds:
        j = db.query(Journal).filter(Journal.id == s["id"]).first()
        if not j:
            j = Journal(
                id=s["id"],
                name=s["name"],
                publisher=s["publisher"],
                scope=s["scope"],
                citation_style=s["citation_style"],
                official_url=s["official_url"],
                is_active=True,
            )
            db.add(j)
            db.flush()

            gv = GuidelineVersion(
                journal_id=j.id,
                version_tag=s["version_tag"],
                provenance=s["provenance"],
                is_current=True,
            )
            db.add(gv)
            db.flush()

            for r in s.get("rules", []):
                jr = JournalRule(
                    guideline_version_id=gv.id,
                    rule_code=r["rule_code"],
                    category=r["category"],
                    requirement_description=r["requirement_description"],
                    rule_parameters=r.get("rule_parameters"),
                    severity=r.get("severity", "major"),
                )
                db.add(jr)

            for c in s.get("chunks", []):
                gc = GuidelineChunk(
                    guideline_version_id=gv.id,
                    chunk_index=c["chunk_index"],
                    section_heading=c.get("section_heading"),
                    content=c["content"],
                    source_url=c.get("source_url"),
                    embedding=c.get("embedding"),
                )
                db.add(gc)

    db.commit()
    db.close()

    yield

    Base.metadata.drop_all(bind=test_engine)
    if os.path.exists("./test_platform.db"):
        try:
            os.remove("./test_platform.db")
        except Exception:
            pass


@pytest.fixture
def db_session():
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def valid_pdf_bytes() -> bytes:
    """Generates a synthetic well-structured research paper PDF with citations and figures."""
    doc = fitz.open()

    # Page 1: Title, Abstract, Introduction
    p1 = doc.new_page()
    p1.insert_text((50, 50), "Deep Verification of Academic Manuscripts for Journal Readiness", fontsize=14)
    p1.insert_textbox(
        fitz.Rect(50, 70, 550, 180),
        "Abstract\nAutomated manuscript review requires multi-stage verification combining deterministic checks, "
        "vector guideline retrieval, and LLM academic quality evaluation. In this paper, we propose a modular architecture "
        "that accurately calculates an explainable readiness score from 0 to 100 with actionable feedback. Our experiments "
        "demonstrate 99.4% precision in detecting missing citations and section compliance on diverse benchmarks.",
        fontsize=10,
    )
    p1.insert_textbox(
        fitz.Rect(50, 190, 550, 320),
        "1. Introduction\nAcademic peer review has faced growing submission volumes [1]. Authors frequently encounter formatting\n"
        "rejections due to non-compliant references [2] or missing statements. Prior art [1] explored basic keyword matching.\n"
        "As shown in Figure 1, our proposed framework decouples deterministic verification from semantic heuristics.",
        fontsize=10,
    )
    p1.insert_text((50, 340), "Figure 1: Architectural pipeline of the multi-engine verification platform.", fontsize=9)

    # Page 2: Methodology & Results
    p2 = doc.new_page()
    p2.insert_textbox(
        fitz.Rect(50, 50, 550, 160),
        "2. Methodology\nWe formally specify rule checkers for citations, sections, and layout. In-text citations are mapped\n"
        "to bibliography indices. For RAG guidelines, official publisher rules are isolated by journal metadata to prevent\n"
        "rule contamination. Table 1 summarizes the benchmark datasets.",
        fontsize=10,
    )
    p2.insert_text((50, 180), "Table 1: Evaluation dataset characteristics across publishers.", fontsize=9)
    p2.insert_textbox(
        fitz.Rect(50, 200, 550, 300),
        "3. Experimental Results\nEmpirical evaluations demonstrate that our system processes a 10-page manuscript in under 4 seconds.\n"
        "Statistical tests confirm significant improvements with p-value < 0.001 across baselines.",
        fontsize=10,
    )
    p2.insert_textbox(
        fitz.Rect(50, 310, 550, 390),
        "4. Conclusion\nWe presented a robust manuscript verification engine. Future work will investigate automated camera-ready formatting.",
        fontsize=10,
    )
    p2.insert_textbox(
        fitz.Rect(50, 400, 550, 520),
        "References\n"
        "[1] A. Smith and B. Doe, \"Automated Analysis of Scientific Literature,\" IEEE Trans. Knowl. Data Eng., vol. 34, no. 2, pp. 110-120, 2022.\n"
        "[2] C. Johnson, \"Citation Reliability in Machine Learning Venues,\" Nature Mach. Intell., vol. 4, pp. 45-56, 2023.\n",
        fontsize=9,
    )

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def flawed_pdf_bytes() -> bytes:
    """
    Generates a PDF manuscript with known deliberate errors:
    - Missing citation reference: cites [99] which is not in bibliography
    - Uncited reference: bibliography has [3], never cited in body
    - Missing Conclusion section
    - Unreferenced Figure 2
    """
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((50, 60), "A Study on Unchecked Verification Anomalies", fontsize=14)
    p1.insert_text((50, 100), "Abstract\nShort abstract.", fontsize=10)  # Abstract too short! (< 50 words)
    p1.insert_text((50, 180), "1. Introduction\nHere we cite [1] and non-existent reference [99].", fontsize=10)
    p1.insert_text((50, 240), "Figure 2: An unreferenced figure that is never mentioned in text.", fontsize=9)
    p1.insert_text((50, 300), "2. Methodology\nSimple discussion without conclusion.", fontsize=10)
    p1.insert_text(
        (50, 400),
        "References\n"
        "[1] J. Doe, \"First Paper,\" Journal of Science, 2021.\n"
        "[3] M. Uncited, \"Uncited Reference Entry,\" Tech Report, 2020.\n",
        fontsize=9,
    )

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def scanned_pdf_bytes() -> bytes:
    """Generates a synthetic scanned PDF (raster image page with zero text)."""
    doc = fitz.open()
    page = doc.new_page()
    # Draw a rectangle image simulating a scanned document page
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 200), False)
    pix.clear_with(240)
    page.insert_image(fitz.Rect(50, 50, 350, 350), pixmap=pix)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def valid_docx_bytes() -> bytes:
    """Generates a valid test DOCX manuscript."""
    doc = DocxDoc()
    doc.add_heading("Automated Verification in Word Documents", level=0)
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This is a comprehensive abstract describing the docx verification procedure. " * 5)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Introduction discussing citation [1] and citation [2]. As shown in Figure 1, the pipeline works well.")
    doc.add_paragraph("Figure 1: Pipeline overview in DOCX.")
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("Methodology section describing algorithms and empirical setups.")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Experimental results showing high accuracy.")
    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("Concluding summary of contributions.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] X. Author, 'DOCX Processing,' 2023.")
    doc.add_paragraph("[2] Y. Author, 'Manuscript Ingestion,' 2024.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
